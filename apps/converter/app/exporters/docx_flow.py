"""PDF -> editable Word document.

Paragraphs, tables and images are rebuilt by pdf2docx (with our patches), then
repaired for Khmer (complex-script fonts, sizes, language) and fonts are embedded.
Scanned or broken pages are rebuilt from Khmer OCR.
"""

from __future__ import annotations

import logging
import os
import tempfile

import fitz
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt

from .. import fonts
from ..pipeline import analyze, ocr, textfix
from . import pdf2docx_patches
from .common import Progress
from .docx_common import anchor_picture, font_measurer, is_khmer_capable, polish_runs, render_clip, set_document_defaults

log = logging.getLogger(__name__)


def _offpage_images(doc: fitz.Document) -> dict[int, list[tuple[int, fitz.Rect]]]:
    """Images that stick out of the page (e.g. a flag in the corner). pdf2docx lets them
    push the text around and then loses them, so they are handled separately."""
    found: dict[int, list[tuple[int, fitz.Rect]]] = {}
    for page in doc:
        area = page.rect + (-1, -1, 1, 1)
        for info in page.get_image_info(xrefs=True):
            r = fitz.Rect(info["bbox"])
            if info.get("xref") and not area.contains(r) and abs(r & page.rect) > 25:
                found.setdefault(page.number, []).append((info["xref"], r))
    return found


def _restore_images(docx, pdf_path: str, images: dict[int, list[tuple[int, fitz.Rect]]], n_pages: int) -> None:
    body = docx.element.body
    # paragraphs grouped by section (= PDF page); a section ends with a paragraph holding w:sectPr
    sections: list[list] = [[]]
    for child in body:
        if child.tag == qn("w:p"):
            sections[-1].append(child)
            if child.find(qn("w:pPr") + "/" + qn("w:sectPr")) is not None:
                sections.append([])
    if not sections[-1]:
        sections.pop()
    if len(sections) != n_pages:
        log.info("section/page mismatch (%d vs %d); off-page images not restored", len(sections), n_pages)
        return
    from docx.text.paragraph import Paragraph

    src = fitz.open(pdf_path)
    for pno, items in images.items():
        if not sections[pno]:
            continue
        para = Paragraph(sections[pno][0], docx.part)
        for _, rect in items:
            clip = rect & src[pno].rect
            png = render_clip(src[pno], clip, 200)
            anchor_picture(para.add_run(), png, clip.width, clip.height, clip.x0, clip.y0)
    src.close()


def _pdf2docx(pdf_path: str, out_path: str, pages: list[int], password: str | None) -> None:
    from pdf2docx import Converter

    pdf2docx_patches.apply()
    cv = Converter(pdf_path, password=password) if password else Converter(pdf_path)
    try:
        cv.convert(out_path, pages=pages, multi_processing=False)
    finally:
        cv.close()


def _ocr_docx(doc_pdf: fitz.Document, pages: list[int], out_path: str, progress: Progress | None) -> None:
    d = Document()
    body = d.element.body
    for p in list(body):
        if p.tag == qn("w:p"):
            body.remove(p)
    font = fonts.default_khmer_font()
    for n, idx in enumerate(pages):
        page = doc_pdf[idx]
        if progress:
            progress(n / max(len(pages), 1), f"OCR page {idx + 1}")
        blocks = ocr.ocr_page(page)
        section = d.sections[0] if n == 0 else d.add_section(WD_SECTION.NEW_PAGE)
        W, H = page.rect.width, page.rect.height
        section.page_width, section.page_height = Pt(W), Pt(H)
        if blocks:
            left = max(18.0, min(b.bbox[0] for b in blocks) - 2)
            right = max(18.0, W - max(b.bbox[2] for b in blocks) - 2)
            top = max(18.0, min(b.bbox[1] for b in blocks) - 4)
        else:
            left = right = top = 54.0
        section.left_margin, section.right_margin = Pt(left), Pt(right)
        section.top_margin, section.bottom_margin = Pt(top), Pt(12)
        content_w = W - left - right

        # lay out with the measured OCR line pitch (Word's own "single" Khmer line is far taller),
        # then shrink everything if it would still not fit on one page
        plan = []
        cursor = top
        measure = font_measurer(font)
        for b in blocks:
            if not b.text.strip():
                continue
            heights = sorted(l.height for l in b.lines)
            size = ocr.estimate_font_size(heights[len(heights) // 2], b.text)
            if measure:
                # shrink until the longest line fits where it sits, so Word does not wrap it
                avail = max(40.0, W - right - max(b.bbox[0], left)) * 0.98
                widest = max((measure(ln.text.strip(), size) for ln in b.lines), default=0)
                if widest > avail:
                    size = max(6.0, size * avail / widest)
            n = len(b.lines)
            pitch = (b.lines[-1].bbox[1] - b.lines[0].bbox[1]) / (n - 1) if n > 1 else heights[-1] * 1.3
            pitch = max(pitch, size * 1.2)
            before = max(0.0, b.bbox[1] - cursor)
            plan.append((b, size, pitch, before))
            cursor = max(cursor, b.bbox[1]) + n * pitch
        used = cursor - top
        room = H - top - 14
        scale = min(1.0, room / used) if used > 0 else 1.0

        for b, size, pitch, before in plan:
            # keep the page's own line breaks so lines don't reflow differently from the PDF
            text = "\n".join(ln.text.strip() for ln in b.lines if ln.text.strip())
            size = max(6.0, round(size * scale * 2) / 2)
            para = d.add_paragraph()
            pf = para.paragraph_format
            pf.space_before = Pt(round(before * scale, 1))
            pf.space_after = Pt(0)
            pf.line_spacing = Pt(round(pitch * scale, 1))
            bw = b.bbox[2] - b.bbox[0]
            center = (b.bbox[0] + b.bbox[2]) / 2
            if abs(center - (left + content_w / 2)) < 0.06 * content_w and bw < 0.8 * content_w:
                pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif b.bbox[0] - left > 0.45 * content_w:
                pf.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            else:
                indent = b.bbox[0] - left
                if indent > 8:
                    pf.left_indent = Pt(indent)
            run = para.add_run(text)
            run.font.size = Pt(size)
            run.font.name = font
    d.save(out_path)


def convert(pdf_path: str, out_path: str, infos: list[analyze.PageInfo], options: dict,
                 progress: Progress | None = None) -> list[str]:
    warnings: list[str] = []
    ocr_mode = options.get("ocr", "auto")
    password = options.get("password")
    use_ocr = ocr_mode != "off" and ocr.available()
    if ocr_mode != "off" and not ocr.available() and any(i.needs_ocr for i in infos):
        warnings.append("OCR engine is not installed; scanned pages are kept as images.")

    routes = []
    for info in infos:
        if use_ocr and (ocr_mode == "force" or info.needs_ocr):
            routes.append("ocr")
        else:
            routes.append("layout")
    groups: list[tuple[str, list[int]]] = []
    for i, r in enumerate(routes):
        if groups and groups[-1][0] == r:
            groups[-1][1].append(i)
        else:
            groups.append((r, [i]))

    orders = [i.order for i in infos if i.kind == "text" and i.order != "unknown"]
    order = max(set(orders), key=orders.count) if orders else "unknown"

    src = fitz.open(pdf_path)
    if src.needs_pass:
        src.authenticate(password or "")
    tmpdir = tempfile.mkdtemp(prefix="docx_")
    floating = _offpage_images(src)
    layout_pdf = pdf_path
    if floating:
        hidden = fitz.open(pdf_path)
        if hidden.needs_pass:
            hidden.authenticate(password or "")
        for pno, items in floating.items():
            page = hidden[pno]
            others = [fitz.Rect(i["bbox"]) for i in page.get_image_info(xrefs=True)
                      if all(i["xref"] != x for x, _ in items)]
            for xref, rect in items:
                clip = rect & page.rect
                if any(abs(clip & o) > 0.1 * abs(o) for o in others if abs(o)):
                    continue  # would also erase a normal picture under it
                # remove the picture itself (an emptied placeholder would still distort the layout)
                page.add_redact_annot(clip, fill=False)
            try:
                kwargs = {"images": fitz.PDF_REDACT_IMAGE_REMOVE}
                if hasattr(fitz, "PDF_REDACT_LINE_ART_NONE"):
                    kwargs["graphics"] = fitz.PDF_REDACT_LINE_ART_NONE
                if hasattr(fitz, "PDF_REDACT_TEXT_NONE"):
                    kwargs["text"] = fitz.PDF_REDACT_TEXT_NONE
                page.apply_redactions(**kwargs)
            except Exception:
                log.warning("could not hide off-page images on page %s", pno + 1)
        layout_pdf = os.path.join(tmpdir, "layout.pdf")
        hidden.save(layout_pdf, encryption=fitz.PDF_ENCRYPT_NONE)
        hidden.close()
    parts: list[tuple[str, str]] = []
    done = 0
    for gi, (route, pages) in enumerate(groups):
        part = os.path.join(tmpdir, f"part{gi}.docx")

        def sub(frac: float, msg: str, _done=done, _n=len(pages)) -> None:
            if progress:
                progress((_done + frac * _n) / len(infos), msg)

        if route == "ocr":
            _ocr_docx(src, pages, part, sub)
        else:
            sub(0, f"Layout pages {pages[0] + 1}-{pages[-1] + 1}")
            _pdf2docx(layout_pdf, part, pages, None if floating else password)
        parts.append((route, part))
        done += len(pages)
    src.close()

    if progress:
        progress(0.95, "Finishing document")
    master = Document(parts[0][1])
    if len(parts) > 1:
        from docxcompose.composer import Composer

        composer = Composer(master)
        for _, p in parts[1:]:
            master.add_section(WD_SECTION.NEW_PAGE)
            composer.append(Document(p))
        master = composer.doc

    if floating:
        _restore_images(master, pdf_path, floating, len(infos))
    # every PDF page is its own section with a page break, so a large bottom margin only makes
    # pages spill over when Word's text is a hair taller than the PDF's
    for section in master.sections:
        if section.bottom_margin is None or section.bottom_margin > Pt(12):
            section.bottom_margin = Pt(12)

    # text was already fixed while pdf2docx read it; only fonts/properties are polished here
    used = polish_runs(master, order, fix_text=textfix.current() is None)
    kfont = next((f for f in used if is_khmer_capable(f)), fonts.default_khmer_font())
    set_document_defaults(master, kfont)
    master.save(out_path)

    if options.get("embedFonts", True):
        embedded = fonts.embed_fonts_in_docx(out_path, {f for f in used if is_khmer_capable(f)})
        if embedded:
            log.info("embedded fonts: %s", embedded)
    missing = sorted(f for f in used if is_khmer_capable(f) and not fonts.find_family(f))
    if missing:
        warnings.append("Fonts not available on the server (not embedded): " + ", ".join(missing))
    return warnings


# ---------------------------------------------------------------------------
# Exact mode
# ---------------------------------------------------------------------------
