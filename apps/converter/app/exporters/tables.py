"""PDF -> Excel / CSV: ruled tables (with merged cells and fills), or the text
layout row by row when a document has no tables."""

from __future__ import annotations

import csv
import io
import os
import zipfile
from dataclasses import dataclass, field

import fitz
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .. import fonts, khmer
from ..khmer import symbols
from ..pipeline import analyze, layout, ocr
from .common import Progress, line_text, open_pdf


@dataclass
class Grid:
    page: int
    rows: list[list[str]]
    merges: list[tuple[int, int, int, int]] = field(default_factory=list)  # r0, c0, r1, c1 (0-based, inclusive)
    header_rows: int = 0
    fills: dict[tuple[int, int], str] = field(default_factory=dict)
    is_table: bool = True


def _cluster(values: list[float], tol: float = 2.0) -> list[float]:
    out: list[float] = []
    for v in sorted(values):
        if not out or v - out[-1] > tol:
            out.append(v)
    return out


def _index(edges: list[float], v: float) -> int:
    return min(range(len(edges)), key=lambda i: abs(edges[i] - v))


def _fill_hex(page: fitz.Page, rect: fitz.Rect, drawings) -> str | None:
    for d in drawings:
        fill = d.get("fill")
        if not fill or fill == (1, 1, 1):
            continue
        r = d["rect"]
        if r.contains(rect.tl + (1, 1)) and r.contains(rect.br - (1, 1)):
            return "".join(f"{int(c * 255):02X}" for c in fill[:3])
    return None


def _tables_on_page(page: fitz.Page, info: analyze.PageInfo) -> list[Grid]:
    return [g for g, _ in tables_with_rects(page, info)]


def tables_with_rects(page: fitz.Page, info: analyze.PageInfo) -> list[tuple[Grid, fitz.Rect]]:
    grids: list[tuple[Grid, fitz.Rect]] = []
    try:
        found = page.find_tables()
    except Exception:
        return grids
    try:
        drawings = page.get_drawings()
    except Exception:
        drawings = []
    page_lines = layout.extract_lines(page, info.order) if found.tables else []
    for tab in found.tables:
        cells = [c for c in tab.cells if c]
        if not cells:
            continue
        xs = _cluster([c[0] for c in cells] + [c[2] for c in cells])
        ys = _cluster([c[1] for c in cells] + [c[3] for c in cells])
        nr, nc = max(len(ys) - 1, 1), max(len(xs) - 1, 1)
        rows = [["" for _ in range(nc)] for _ in range(nr)]
        merges: list[tuple[int, int, int, int]] = []
        fills: dict[tuple[int, int], str] = {}
        for c in cells:
            rect = fitz.Rect(c)
            r0, r1 = _index(ys, c[1]), _index(ys, c[3]) - 1
            c0, c1 = _index(xs, c[0]), _index(xs, c[2]) - 1
            r1, c1 = max(r0, r1), max(c0, c1)
            text = _text_in_rect(page_lines, rect)
            if r0 < nr and c0 < nc:
                rows[r0][c0] = text
                if r1 > r0 or c1 > c0:
                    merges.append((r0, c0, min(r1, nr - 1), min(c1, nc - 1)))
                fill = _fill_hex(page, rect, drawings)
                if fill:
                    fills[(r0, c0)] = fill
        header = 1 if getattr(tab, "header", None) and not tab.header.external else 0
        grids.append((Grid(page.number, rows, merges, header, fills), fitz.Rect(tab.bbox)))
    return grids


def _text_in_rect(lines: list[layout.Line], rect: fitz.Rect) -> str:
    out = ""
    for ln in lines:
        x0, y0, x1, y1 = ln.bbox
        cx, cy = (x0 + x1) / 2, (y0 + y1) / 2
        if rect.x0 <= cx <= rect.x1 and rect.y0 <= cy <= rect.y1:
            out = khmer.join_khmer_lines(out, line_text(ln))
    return out.strip()


def _lines_to_rows(lines: list[tuple[tuple[float, float, float, float], list[tuple[float, float, str]]]]) -> list[list[str]]:
    """Text without a ruled table: split each line into columns on wide gaps."""
    rows = []
    for _, words in lines:
        cols: list[str] = []
        cur = ""
        prev_x1 = None
        for x0, x1, w in words:
            if prev_x1 is not None and x0 - prev_x1 > 14:
                cols.append(cur)
                cur = w
            else:
                cur = khmer.join_khmer_lines(cur, w) if cur else w
            prev_x1 = x1
        cols.append(cur)
        rows.append(cols)
    return rows


def _text_grid(page: fitz.Page, info: analyze.PageInfo, use_ocr: bool) -> Grid:
    lines = []
    if info.needs_ocr and use_ocr:
        for b in ocr.ocr_page(page):
            for ln in b.lines:
                lines.append((ln.bbox, [(ln.bbox[0], ln.bbox[2], ln.text)]))
    else:
        for ln in layout.extract_lines(page, info.order):
            lines.append((ln.bbox, [(s.bbox[0], s.bbox[2], symbols.to_unicode(s.text, s.font).strip())
                                    for s in ln.spans if s.text.strip()]))
    lines.sort(key=lambda l: (round(l[0][1] / 3), l[0][0]))
    # merge lines that sit on the same baseline
    merged: list = []
    for bbox, words in lines:
        if merged and abs(merged[-1][0][1] - bbox[1]) < 3:
            merged[-1] = (merged[-1][0], sorted(merged[-1][1] + words))
        else:
            merged.append((bbox, sorted(words)))
    return Grid(page.number, _lines_to_rows(merged), is_table=False)


def extract_grids(pdf_path: str, infos: list[analyze.PageInfo], options: dict, progress: Progress | None) -> list[Grid]:
    use_ocr = options.get("ocr", "auto") != "off" and ocr.available()
    doc = open_pdf(pdf_path, options)
    grids: list[Grid] = []
    for n, page in enumerate(doc):
        if progress:
            progress(n / len(doc), f"Page {n + 1}")
        info = infos[n]
        tables = [] if info.needs_ocr else _tables_on_page(page, info)
        grids.extend(tables)
    if not grids:  # no ruled tables anywhere -> export the text layout row by row
        for n, page in enumerate(doc):
            g = _text_grid(page, infos[n], use_ocr)
            if g.rows:
                grids.append(g)
    doc.close()
    return grids


def _dominant_font(pdf_path: str, options: dict) -> str:
    doc = open_pdf(pdf_path, options)
    counts: dict[str, int] = {}
    for page in list(doc)[:5]:
        for ln in layout.extract_lines(page):
            for s in ln.spans:
                counts[s.font] = counts.get(s.font, 0) + len(s.text)
    doc.close()
    khmer_fonts = {k: v for k, v in counts.items() if any(h in k.lower() for h in ("khmer", "battambang", "siemreap", "moul", "kantumruy", "hanuman", "mef"))}
    pool = khmer_fonts or counts
    return max(pool, key=pool.get) if pool else fonts.default_khmer_font()


def to_xlsx(pdf_path: str, out_path: str, infos, options: dict, progress: Progress | None = None) -> list[str]:
    grids = extract_grids(pdf_path, infos, options, progress)
    font_name = _dominant_font(pdf_path, options)
    wb = Workbook()
    wb.remove(wb.active)
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    if not grids:
        wb.create_sheet("Sheet1")
    counters: dict[int, int] = {}
    for g in grids:
        counters[g.page] = counters.get(g.page, 0) + 1
        title = f"P{g.page + 1}" + (f"-T{counters[g.page]}" if g.is_table else "")
        ws = wb.create_sheet(title[:31])
        widths: dict[int, float] = {}
        for r, row in enumerate(g.rows, start=1):
            for c, val in enumerate(row, start=1):
                cell = ws.cell(row=r, column=c, value=_typed(val))
                cell.font = Font(name=font_name, size=11, bold=r <= g.header_rows)
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                if g.is_table:
                    cell.border = border
                fill = g.fills.get((r - 1, c - 1))
                if fill:
                    cell.fill = PatternFill("solid", fgColor=fill)
                longest = max((len(x) for x in str(val).split("\n")), default=0)
                widths[c] = max(widths.get(c, 8), min(60, longest * (1.3 if khmer.contains_khmer(str(val)) else 1.1) + 2))
        for r0, c0, r1, c1 in g.merges:
            ws.merge_cells(start_row=r0 + 1, start_column=c0 + 1, end_row=r1 + 1, end_column=c1 + 1)
        for c, w in widths.items():
            ws.column_dimensions[get_column_letter(c)].width = w
        if g.header_rows:
            ws.freeze_panes = ws.cell(row=g.header_rows + 1, column=1)
    wb.save(out_path)
    return [] if any(g.is_table for g in grids) else ["No ruled tables found; text was exported line by line."]


def _typed(val: str):
    v = (val or "").strip()
    if not v:
        return None
    plain = v.replace(",", "")
    try:
        if plain.isdigit() and not (len(plain) > 1 and plain.startswith("0")) and len(plain) < 15:
            return int(plain)
        if plain.count(".") == 1 and plain.replace(".", "").isdigit():
            return float(plain)
    except ValueError:
        pass
    return v


def to_csv(pdf_path: str, out_dir: str, base: str, infos, options: dict, progress: Progress | None = None) -> tuple[str, list[str]]:
    grids = extract_grids(pdf_path, infos, options, progress)
    warnings = [] if any(g.is_table for g in grids) else ["No ruled tables found; text was exported line by line."]

    def write(g: Grid) -> bytes:
        buf = io.StringIO()
        w = csv.writer(buf)
        for row in g.rows:
            w.writerow(row)
        # BOM so Excel opens Khmer UTF-8 correctly
        return ("﻿" + buf.getvalue()).encode("utf-8")

    if len(grids) <= 1:
        path = os.path.join(out_dir, base + ".csv")
        with open(path, "wb") as f:
            f.write(write(grids[0]) if grids else "﻿".encode())
        return path, warnings
    path = os.path.join(out_dir, base + "-csv.zip")
    counters: dict[int, int] = {}
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for g in grids:
            counters[g.page] = counters.get(g.page, 0) + 1
            z.writestr(f"{base}-page{g.page + 1}-table{counters[g.page]}.csv", write(g))
    return path, warnings
